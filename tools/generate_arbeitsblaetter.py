"""Build the reviewed A5 worksheets from the shared content source (python-docx)."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'materialien/arbeitsblaetter/arbeitsblaetter_inhalte.json'
OUT=ROOT/'materialien/arbeitsblaetter/Arbeitsblaetter_1-24_DIN_A5_ueberarbeitet.docx'
ILLUSTRATION=ROOT/'materialien/Collage Otterbilder Briefe.png'
# Native Word crops preserve the supplied artwork without making new images.
# Bounding boxes are fractions of the original collage, not stretched tiles.
MOTIFS={
 'denken':(.368,.31,.602,.603),
 'lesen':(.055,.335,.328,.641),
 'schreiben':(.345,.037,.635,.298),
 'post':(.029,.009,.331,.321),
 'fertig':(.654,.611,.946,.95),
}
MOTIF_BY_SHEET=['denken','denken','denken','post','post','lesen',
 'lesen','lesen','schreiben','lesen','schreiben','schreiben',
 'denken','denken','denken','denken','schreiben','schreiben',
 'lesen','schreiben','schreiben','fertig','denken','fertig']

def add_otter(paragraph, number):
 """Place one cropped motif in the unused top-right header space."""
 motif=MOTIF_BY_SHEET[number-1]
 left,top,right,bottom=MOTIFS[motif]
 # Source aspect ratio is 1448:1086. Keep the crop's natural proportions.
 height=Cm(1.55)
 width=round(height*(1448/1086)*(right-left)/(bottom-top))
 picture=paragraph.add_run().add_picture(str(ILLUSTRATION),width=width,height=height)
 inline=picture._inline
 crop=OxmlElement('a:srcRect')
 for key,value in [('l',left),('t',top),('r',1-right),('b',1-bottom)]:
  crop.set(key,str(round(value*100000)))
 fill=inline.find('.//'+qn('pic:blipFill'))
 fill.insert(1,crop)
 if motif=='lesen':
  # The next collage motif starts beneath the chair. Clip only that corner
  # with Word's native picture geometry; retain the reader's complete tail.
  props=inline.find('.//'+qn('pic:spPr'))
  shape=props.find(qn('a:prstGeom'))
  geom=OxmlElement('a:custGeom')
  for name in ['avLst','gdLst','ahLst','cxnLst']:geom.append(OxmlElement('a:'+name))
  rect=OxmlElement('a:rect')
  for key,value in [('l','0'),('t','0'),('r','r'),('b','b')]:rect.set(key,value)
  geom.append(rect)
  paths=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w','100000');path.set('h','100000')
  for i,(x,y) in enumerate([(0,0),(100000,0),(100000,94500),(44000,94500),(44000,100000),(0,100000)]):
   command=OxmlElement('a:moveTo' if i==0 else 'a:lnTo');point=OxmlElement('a:pt');point.set('x',str(x));point.set('y',str(y));command.append(point);path.append(command)
  path.append(OxmlElement('a:close'));paths.append(path);geom.append(paths)
  props.replace(shape,geom)
 inline.docPr.set('descr','Otter '+{'denken':'denkt nach','lesen':'liest einen Brief','schreiben':'schreibt einen Brief','post':'bringt einen Brief','fertig':'zeigt einen fertigen Brief'}[motif])
 anchor=OxmlElement('wp:anchor')
 for key,value in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'0'}.items():
  anchor.set(key,value)
 pos=OxmlElement('wp:simplePos');pos.set('x','0');pos.set('y','0');anchor.append(pos)
 horizontal=OxmlElement('wp:positionH');horizontal.set('relativeFrom','column')
 align=OxmlElement('wp:align');align.text='right';horizontal.append(align);anchor.append(horizontal)
 vertical=OxmlElement('wp:positionV');vertical.set('relativeFrom','paragraph')
 offset=OxmlElement('wp:posOffset');offset.text='0';vertical.append(offset);anchor.append(vertical)
 for tag in ['extent','effectExtent']:
  child=inline.find(qn('wp:'+tag))
  if child is not None:anchor.append(child)
 anchor.append(OxmlElement('wp:wrapNone'))
 for child in list(inline):anchor.append(child)
 inline.getparent().replace(inline,anchor)
 paragraph.paragraph_format.right_indent=Cm(3)
D=Document()
s=D.sections[0];s.page_width=Cm(14.8);s.page_height=Cm(21)
s.top_margin=Cm(.8);s.bottom_margin=Cm(.8);s.left_margin=Cm(.9);s.right_margin=Cm(.9)
s.footer_distance=Cm(.3)
for name,size in [('Normal',14),('Heading 1',16),('Heading 2',14)]:
 st=D.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('24313A')
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.space_before=Pt(0);st.paragraph_format.line_spacing=1.02
 st.paragraph_format.keep_with_next=name!='Normal'
D.core_properties.title='Wunschbriefe · Arbeitsblätter 1–24 · Fassung 3'
D.core_properties.author='Deutschunterricht'
D.core_properties.subject='Schritt 2 · 7. September 2026'
def para(t,bold=False,size=None):
 p=D.add_paragraph();r=p.add_run(t);r.bold=bold
 if size:r.font.size=Pt(size)
 return p

def band(label,t,color='EEF2F4'):
 p=D.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(6)
 r=p.add_run(label+('\n' if '\n' in t else ': '));r.bold=True;p.add_run(t)
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);p._p.get_or_add_pPr().append(sh)
 return p

def table(rows):
 widths=[1.5,11.5] if rows[0][0]=='Nr.' else [8.1,4.9]
 ts=[round(x*1440/2.54) for x in widths]
 t=D.add_table(rows=0,cols=2);t.autofit=False
 pr=t._tbl.tblPr
 for name,value in [('tblW',sum(ts)),('tblInd',80)]:
  el=pr.find(qn('w:'+name))
  if el is None:el=OxmlElement('w:'+name);pr.append(el)
  el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa')
 m=OxmlElement('w:tblCellMar')
 for k in ['top','bottom','left','right']:
  el=OxmlElement('w:'+k);el.set(qn('w:w'),'60' if k in ['top','bottom'] else '80');el.set(qn('w:type'),'dxa');m.append(el)
 pr.append(m)
 for i,v in enumerate(ts):t._tbl.tblGrid.gridCol_lst[i].set(qn('w:w'),str(v))
 for rowi,row in enumerate(rows):
  rw=t.add_row();rw._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,txt in enumerate(row):
   c=rw.cells[j];c.width=Cm(widths[j]);c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(ts[j]))
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);r=p.add_run(txt);r.bold=rowi==0
   if rowi==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF2');c._tc.get_or_add_tcPr().append(sh)
  if rowi==0:rw._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 return t
for item in json.loads(DATA.read_text()):
 p=para('Blatt '+str(item['n']),True,20);p.paragraph_format.space_after=Pt(3);p.paragraph_format.page_break_before=item['n']>1
 add_otter(p,item['n'])
 para(item['status'],True).paragraph_format.right_indent=Cm(3)
 D.add_paragraph(item['title'],'Heading 1')
 band('Dein Ziel',item['goal'])
 for b in item['blocks']:
  if b[0]=='task':
   p=para(b[1],True);p.paragraph_format.space_before=Pt(12);p.paragraph_format.keep_with_next=True
   para(b[2])
  elif b[0]=='text':para(b[1])
  elif b[0]=='checks':
   for line in b[1]:para('☐ '+line)
  elif b[0]=='box':band(b[1],b[2],'F5F4F1' if b[1]!='Zusatz · freiwillig' else 'F9F0E9')
  elif b[0]=='table':table(b[1])
 band('Hilfe',item['help'],'EDF4EA')
 band('Prüfe',item['check'])
D.save(OUT)
print(OUT)
