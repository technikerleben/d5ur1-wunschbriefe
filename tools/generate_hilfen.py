"""Step 4: eight A5 reference sheets and eight printable strategy strips."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
ROOT=Path(__file__).resolve().parents[1]
D=Document();s=D.sections[0];s.page_width=Cm(14.8);s.page_height=Cm(21)
s.top_margin=s.bottom_margin=Cm(.9);s.left_margin=s.right_margin=Cm(.9)
for name,size in [('Normal',14),('Title',20),('Heading 1',16),('Heading 2',14)]:
 st=D.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('24313A');st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(0);st.paragraph_format.line_spacing=1.04;st.paragraph_format.keep_with_next=name!='Normal'
D.core_properties.title='Merkblätter Wunschbriefe · Fassung 3';D.core_properties.subject='Schritt 4';D.core_properties.author='Deutschunterricht'
def p(t,style=None):return D.add_paragraph(t,style)
def band(label,t):
 x=p('');x.add_run(label+'\n').bold=True;x.add_run(t);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF2F4');x._p.get_or_add_pPr().append(sh)
def sheet(n,title,goal):
 x=p('Merkblatt '+str(n),'Title');x.paragraph_format.page_break_before=n>1
 p(title,'Heading 1');band('Das hilft dir:',goal)
sheet(1,'Eine Schreibsituation verstehen','Du klärst, an wen du schreibst und was du möchtest.')
p('Lies den Auftrag. Kläre diese Fragen:','Heading 2')
for t in ['Wer schreibt den Brief?','Wer bekommt den Brief?','Was wünschst du dir für die Schule?','Welcher weitere Satz passt dazu?']:p(t)
band('Beispiel','Du schreibst an die Schulleitung. Du wünschst dir Sitzbänke. Du ergänzt: „Ich würde mich darüber sehr freuen.“')
p('Sage deinen Plan erst mündlich. Notiere dann Stichwörter. Blatt 16 hilft dir beim Planen.')
band('Zusatz · freiwillig','Du kannst einen Grund oder ein passendes Umfrageergebnis ergänzen. Ohne diese Zusätze kann dein Brief auch gelingen.')
sheet(2,'So ist ein Brief aufgebaut','Du ordnest die fünf Briefteile.')
for t in ['1  Ort und Datum','2  Anrede','3  Brieftext: Wunsch und weiterer Satz','4  Grußformel','5  Name']:p(t)
band('Beispiel','Dortmund, 15. September 2026\nSehr geehrte Frau Berger,\nfür unsere Klasse wünsche ich mir eine Leseecke. Ich würde mich darüber sehr freuen.\nFreundliche Grüße\nMina')
p('Nach der Anrede steht ein Komma. Nach der Grußformel steht kein Komma. Beginne jeden Briefteil in einer neuen Zeile.')
sheet(3,'Wünsche freundlich formulieren','Du schreibst freundlich und genau.')
p('Wähle einen passenden Anfang:','Heading 2')
for t in ['Ich wünsche mir …','Für unsere Klasse wünsche ich mir …','Ich würde mich freuen, wenn …','Könnten Sie bitte …?']:p(t)
band('Freundlich und genau','Statt: „Geben Sie uns mehr Sachen!“\nSchreibe: „Ich wünsche mir neue Bälle für die Pause.“')
p('Ein weiterer passender Satz','Heading 2')
p('Du kannst sagen, was du erhoffst oder wie du die Sache nutzen möchtest.')
band('Beispiele','Ich würde mich darüber sehr freuen.\nDie Bälle könnten wir in der Pause nutzen.')
p('Schreibe „Sie“ groß, wenn du eine erwachsene Person höflich ansprichst.')
sheet(4,'Einen Wunsch begründen','Du erklärst einen Nutzen. Diese Vertiefung ist freiwillig.')
band('Ein Grund erklärt den Nutzen','Eine Spieleausleihe wäre sinnvoll, weil wir dann gemeinsam spielen können.')
p('So kannst du beginnen:','Heading 2')
p('Das wäre sinnvoll, weil …\nDas hilft uns, denn …\nDadurch können wir …')
p('„Weil ich das will“ erklärt den Nutzen noch nicht. Frage dich: Was wird durch den Wunsch besser?')
band('Zahlen · ebenfalls freiwillig','18 von 26 Kindern wünschen sich eine Spieleausleihe.\nDas ist ein Umfrageergebnis. Die Zahl allein erklärt noch keinen sachlichen Nutzen.')
p('Übernimm Zahl und Wunsch genau. Schreibe nicht „alle“, wenn es 18 von 26 sind. Nutze nur Daten, die zu deinem Wunsch passen.')
p('Ohne Grund und Zahlen kann dein Grundtext vollständig sein.')
sheet(5,'Einen Brief planen und schreiben','Du gehst Schritt für Schritt vor.')
p('Planung','Heading 1')
p('Lies den Auftrag. Wähle Person und Wunsch. Plane einen weiteren passenden Satz. Notiere Stichwörter mit Blatt 16.')
p('Durchführung','Heading 1')
p('Schreibe ganze Sätze. Nutze alle fünf Briefteile von Merkblatt 2. Blatt 17 hilft dir bei Bedarf mit Satzanfängen.')
p('Reflexion','Heading 1')
p('Prüfe mit Merkblatt 6 oder Blatt 19. Verbessere nur, was noch fehlt oder unklar ist. Ein passender Text darf so bleiben.')
band('Zusatz · freiwillig','Du kannst einen Grund oder passende Daten ergänzen. Prüfe dann auch, ob diese stimmen.')
p('Besprich mit der Lehrkraft, was schon gelingt und welche Übung dir als Nächstes hilft.')
sheet(6,'Checkliste für deinen Wunschbrief','Du prüfst deinen Brief an konkreten Stellen.')
for t in ['Ort und Datum stehen oben.','Die Anrede passt zur Person.','Dein Wunsch ist klar und passt zur Aufgabe.','Mindestens ein weiterer Satz passt dazu.','Dein Brief klingt freundlich.','Grußformel und Name stehen am Ende.','Du schreibst verständliche ganze Sätze.','Deine Schrift ist lesbar.']:p('☐ '+t)
p('Hake erst ab, wenn du die passende Stelle gefunden hast. Lies für Sprache und Lesbarkeit den ganzen Brief.')
band('Nur wenn du Zusätze nutzt','Passt dein Grund zum Wunsch?\nStimmen die Zahlen und der Wunsch aus der Umfrage?')
p('Verbessere nur, was nötig ist. Der Kompetenzcheck hilft dir bei der Wahl deiner nächsten Übung.')
sheet(7,'So arbeitest du in der Lernzeit','Du nutzt Planung, Durchführung und Reflexion.')
p('Planung','Heading 1')
p('Schau auf deinen Papierlernweg. Wähle ein kleines Ziel und dein nächstes Blatt. Lege Heft, Stift und Hilfen bereit.')
p('Durchführung','Heading 1')
p('Bearbeite eine Aufgabe nach der anderen. Der Auftrag sagt dir: auf dem Blatt, im Heft oder mündlich. Nutze Hilfe, wenn du sie brauchst.')
p('Reflexion','Heading 1')
p('Nutze „Prüfe“ auf deinem Blatt. Zeige, was gelungen ist. Hake die bearbeitete und geprüfte Aufgabe im Lernweg ab.')
band('Dein nächster Schritt','Was kannst du schon?\nWas möchtest du noch üben?\nWähle eine passende nächste Aufgabe.')
p('Du musst nicht schnell sein. Viele Haken zeigen erledigte Aufgaben. Dein Ergebnis zeigt, was du kannst.')
sheet(8,'Deine Hilfekette','Du findest Hilfe, wenn du feststeckst.')
p('Lies den Auftrag noch einmal.','Heading 2');p('Schau auf die Hilfe und das Beispiel auf deinem Blatt.')
p('Wähle eine passende Papierhilfe:','Heading 2')
for t in ['Wunsch oder Person? → Merkblatt 1','Briefteile? → Merkblatt 2','Freundliche Sätze? → Merkblatt 3','Grund oder Daten? → Merkblatt 4','Planen oder prüfen? → Merkblatt 5/6']:p(t)
p('Frage ein Partnerkind oder die Lehrkraft. Sage oder zeige, wo du feststeckst.')
band('Du darfst direkt fragen','Kannst du den Auftrag noch nicht lesen oder verstehst du ihn nicht? Bitte sofort um Vorlesen oder Erklären. Du musst nicht alle Hilfen vorher ausprobieren.')
p('Hilfe ist beim Üben erlaubt. Für die Probearbeit und die Arbeit erklärt die Lehrkraft, welche Hilfen du nutzen darfst.')
D.save(ROOT/'materialien/merkblaetter/Merkblaetter_Wunschbriefe_DIN_A5_Mathe_angepasst.docx')

cards=[('PLANUNG','Auftrag klären',['Lies den Auftrag oder lass ihn dir vorlesen.','Markiere wichtige Wörter.','Sage oder zeige, was du tun sollst.','Prüfe: Was soll am Ende entstehen?']),('PLANUNG','Lernziel setzen',['Überlege: Was möchtest du können?','Wähle ein kleines Ziel.','Überlege: Woran erkennst du deinen Erfolg?','Sage dein Ziel oder schreibe es kurz auf.']),('PLANUNG','Arbeit in Schritte teilen',['Schau dir die Aufgabe an.','Teile sie in kleine Schritte.','Lege die Reihenfolge fest.','Beginne mit dem ersten Schritt.','Hake erledigte Schritte ab.']),('DURCHFÜHRUNG','Beispiel nutzen',['Schau dir das Beispiel an.','Markiere, was dir hilft.','Vergleiche es mit deiner Aufgabe.','Nutze den Aufbau mit deinen eigenen Ideen.','Prüfe: Passt dein Ergebnis zur Aufgabe?']),('DURCHFÜHRUNG','Zwischenstopp machen',['Halte kurz an.','Vergleiche deinen Stand mit Auftrag und Ziel.','Prüfe, was schon gelingt.','Wähle deinen nächsten Schritt.','Ändere deinen Plan, wenn es nötig ist.']),('DURCHFÜHRUNG','Hilfekette nutzen',['Lies den Auftrag und die Hilfe auf dem Blatt.','Nutze ein Beispiel oder ein passendes Merkblatt.','Frage ein Partnerkind oder die Lehrkraft.','Zeige oder sage, wo du feststeckst.','Verstehst du den Auftrag nicht? Frage direkt.']),('REFLEXION','Mit Kriterien prüfen',['Lies einen Prüfpunkt nach dem anderen.','Zeige die passende Stelle in deinem Ergebnis.','Prüfe: Ist der Punkt erfüllt?','Verbessere nur, was fehlt oder unklar ist.','Prüfe erneut. Was schon passt, darf bleiben.']),('REFLEXION','Lernprozess reflektieren',['Zeige, was gut gelungen ist.','Nenne eine Hilfe, die dir geholfen hat.','Was war noch schwierig? Sage es, wenn du möchtest.','Wähle deinen nächsten Lernschritt.'])]
file=ROOT/'materialien/strategiekarten/Strategiekarten_Lernbuddy_Querstreifen_8_Karten.pdf'
c=canvas.Canvas(str(file),pagesize=A4);w,h=A4;c.setTitle('Strategiekarten · Fassung 3');c.setAuthor('Deutschunterricht')
for i,(phase,title,steps) in enumerate(cards):
 slot=i%4;top=h-20-slot*200
 c.setFillColorRGB(.14,.19,.23);c.setFont('Helvetica-Bold',12);c.drawString(24,top-12,f'STRATEGIE {i+1} · {phase}')
 c.setFont('Helvetica-Bold',17);c.drawString(24,top-36,title)
 y=top-61
 for k,t in enumerate(steps,1):
  words=t.split();lines=[];line=''
  for word in words:
   trial=(line+' '+word).strip()
   if stringWidth(trial,'Helvetica',14)>w-76:lines.append(line);line=word
   else:line=trial
  if line:lines.append(line)
  c.setFont('Helvetica-Bold',14);c.drawString(24,y,str(k)+'.')
  c.setFont('Helvetica',14)
  for line in lines:c.drawString(46,y,line);y-=17
  y-=5
 assert y>top-193,(i,y,top)
 if slot<3:
  c.setStrokeColorRGB(.5,.5,.5);c.setLineWidth(.5)
  for a,b in [(6,18),(w-18,w-6)]:c.line(a,top-195,b,top-195)
 if slot==3:
  c.setFont('Helvetica',9);c.drawString(24,12,f'Fassung 3 · Seite {i//4+1}/2 · An den Schnittmarken schneiden.');c.showPage()
c.save()
print('Reference DOCX and strategy PDF generated.')
