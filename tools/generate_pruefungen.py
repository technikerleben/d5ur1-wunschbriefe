"""Comparable paper assessments and mock exam using one shared task template."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];EX=ROOT/'materialien/lernerfolgskontrolle';MOCK=ROOT/'materialien/probearbeit'
BASE=EX/'Wunschbrief_Das_zaehlt_Kinderfassung.docx'
helpdoc=Document(BASE);HELP=[deepcopy(x) for x in helpdoc.element.body if x.tag!=qn('w:sectPr')]
for e in HELP:
 for t in e.iter(qn('w:t')):
  if t.text:t.text=t.text.replace('auf Seite 1','unter „Das gehört in deinen Brief“')
# The wording of the existing aid is unchanged except its now-unambiguous cross-reference.
def new(title):
 d=Document(BASE);d.element.body.clear_content()
 for section in d.sections:
  for paragraph in section.footer.paragraphs:
   for run in paragraph.runs:
    if 'September 2026' in run.text or 'Fassung 2' in run.text:
     run.text=run.text.replace('7. September 2026','8. September 2026').replace('Fassung 2','Fassung 3')
 d.core_properties.title=title+' · Fassung 3';d.core_properties.subject='Schritt 7 · 8. September 2026';d.core_properties.author='Deutschunterricht'
 for name,size,bold in [('ExamBody',14,False),('ExamTitle',22,True),('ExamHeading',16,True),('ExamLabel',14,True)]:
  st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string('24313A');st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.03;st.paragraph_format.keep_with_next=name!='ExamBody'
 return d

def p(d,t='',style='ExamBody'):return d.add_paragraph(t,style)
def page(d,title,first=False):
 x=p(d,title,'ExamTitle');x.paragraph_format.page_break_before=not first

def lines(d,n=1,height=26):
 for _ in range(n):
  x=p(d);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=Pt(height)
  x.add_run(' ').font.size=Pt(1)
  borders=OxmlElement('w:pBdr');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'3');b.set(qn('w:color'),'B2B8BC');borders.append(b);between=deepcopy(b);between.tag=qn('w:between');borders.append(between);x._p.get_or_add_pPr().append(borders)
def field(d,label):p(d,label,'ExamLabel');lines(d,1,22)
def table(d,rows):
 widths=[12.8,5.0];tw=[round(x*1440/2.54) for x in widths]
 t=d.add_table(rows=0,cols=2);t.autofit=False;pr=t._tbl.tblPr
 w=pr.find(qn('w:tblW'));w.set(qn('w:w'),str(sum(tw)));w.set(qn('w:type'),'dxa')
 ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'100');ind.set(qn('w:type'),'dxa');pr.append(ind)
 mar=OxmlElement('w:tblCellMar')
 for name in ['top','bottom','left','right']:
  e=OxmlElement('w:'+name);e.set(qn('w:w'),'80' if name in ['top','bottom'] else '100');e.set(qn('w:type'),'dxa');mar.append(e)
 pr.append(mar)
 for i,v in enumerate(tw):t._tbl.tblGrid.gridCol_lst[i].set(qn('w:w'),str(v))
 for i,row in enumerate(rows):
  r=t.add_row();r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,txt in enumerate(row):
   c=r.cells[j];c.width=Cm(widths[j]);x=c.paragraphs[0];x.style=d.styles['ExamBody'];x.paragraph_format.space_after=Pt(2);run=x.add_run(txt);run.bold=i==0
  if i==0:r._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 return t
variants=[('Termin 1','eine Spieleausleihe',['Spieleausleihe','Leseecke','Sitzplätze','Neue Bälle']),('Termin 2','eine Leseecke',['Leseecke','Sitzplätze','Neue Bälle','Spieleausleihe']),('Termin 3','mehr Sitzplätze auf dem Schulhof',['Sitzplätze','Neue Bälle','Spieleausleihe','Leseecke'])]
mock=('Probearbeit','Ablagefächer für unseren Klassenraum',['Ablagefächer','Sitzkissen','Pflanzen','Bilder an der Wand'])
def append_help(d):
 content=deepcopy(HELP)
 # Explicit break before the first aid paragraph; keep the aid's internal page break.
 first=content[0];pr=first.find(qn('w:pPr'))
 if pr is None:pr=OxmlElement('w:pPr');first.insert(0,pr)
 br=OxmlElement('w:pageBreakBefore');pr.append(br)
 for el in content:d.element.body.insert(len(d.element.body)-1,el)
def packet(d,spec,first=False):
 label,wish,items=spec
 page(d,label+' · Dein Wunschbrief',first)
 p(d,'Name: __________________________  Datum: ______________')
 p(d,'Arbeitszeit: ______ Minuten · Die Lehrkraft trägt die Zeit ein.')
 p(d,'Deine Situation','ExamHeading')
 p(d,'Die Schulleitung sammelt Wünsche für den Schulalltag. Schreibe ihr einen Brief. Dein Wunsch ist: '+wish+'.')
 p(d,'Dein Auftrag','ExamHeading')
 for text in ['Schreibe Ort und Datum oben auf den Brief.','Wähle eine passende Anrede für die Schulleitung.','Formuliere den vorgegebenen Wunsch klar und freundlich.','Schreibe mindestens einen weiteren passenden Satz.','Beende den Brief mit einer Grußformel und deinem Namen.','Schreibe ganze Sätze und so, dass man deine Schrift lesen kann.']:p(d,'☐ '+text)
 p(d,'Zusätze · freiwillig','ExamHeading')
 p(d,'Du kannst einen sachlichen Grund und/oder ein passendes Umfrageergebnis ergänzen. Beides darfst du weglassen.')
 p(d,'Beispielumfrage · nur für den freiwilligen Zusatz','ExamHeading')
 p(d,'26 Kinder einer Beispielklasse wurden gefragt: „Was wünschst du dir für den Schulalltag?“ Mehrere Antworten waren erlaubt. Die Zahlen sind erfunden.')
 table(d,[['Wunsch','Stimmen']]+[[x,str(v)] for x,v in zip(items,[17,14,11,8])])
 p(d,'Arbeite mit Planung → Durchführung → Reflexion.')
 if label=='Probearbeit':p(d,'Diese Probearbeit ist verpflichtend und bleibt ohne Note.')
 page(d,label+' · Planung')
 p(d,'Notiere Stichwörter. Deinen Brief schreibst du erst auf den Schreibbogen. Die beiden Hilfeseiten am Ende darfst du nutzen.')
 for label_ in ['An wen schreibst du?','Was wünschst du dir?','Welcher weitere Satz passt zu deinem Wunsch?','Welche Anrede wählst du?','Welche Grußformel wählst du?']:field(d,label_)
 p(d,'Nur wenn du Zusätze nutzen möchtest','ExamHeading')
 for label_ in ['Welchen sachlichen Grund möchtest du nennen? · freiwillig','Welches passende Umfrageergebnis möchtest du nutzen? · freiwillig']:field(d,label_)
 p(d,'Die beiden Zusatzfelder dürfen leer bleiben.')
 page(d,label+' · Durchführung')
 p(d,'Name: __________________________  Datum: ______________')
 p(d,'Schreibe hier deinen vollständigen Brief. Bei Bedarf bekommst du weiteres Schreibpapier.')
 lines(d,21,25)
 p(d,'Reflexion','ExamHeading')
 p(d,'Lies deinen Brief. Prüfe ihn mit „Dein Wunschbrief: Das zählt“. Verbessere nur, was nötig ist. Die zwei Hilfeseiten folgen direkt.')
 append_help(d)

def feedback(d):
 page(d,'Rückmeldung zur Probearbeit')
 p(d,'Name: __________________________  Datum: ______________')
 p(d,'Die Lehrkraft füllt diesen Bogen mit dir aus. Er ist keine Note.')
 p(d,'Betrachteter Text: ☐ erste Fassung  ☐ unterstützte Überarbeitung')
 p(d,'Das zeigt dein Brief','ExamHeading')
 for t in ['Die Anrede passt zur Person.','Dein Wunsch passt zum Auftrag und ist verständlich.','Ein weiterer Satz passt zu deinem Wunsch.','Grußformel und dein Name stehen am Ende.','Dein Brief ist freundlich.','Du schreibst verständliche Sätze und lesbar.']:
  p(d,t);p(d,'☐ gelungen   ☐ noch üben')
 p(d,'Ort und Datum: ☐ vorhanden   ☐ noch ergänzen')
 p(d,'Freiwillige Zusätze','ExamHeading')
 p(d,'Grund: ☐ passend  ☐ noch prüfen  ☐ nicht genutzt\nDaten: ☐ passend  ☐ noch prüfen  ☐ nicht genutzt')
 field(d,'Diese zusätzliche Hilfe gab es (oder: keine):')
 field(d,'Das gelingt dir schon:')
 field(d,'Dein nächster Schritt / passendes Blatt:')
 p(d,'Terminwahl: ☐ besprechen   ☐ nach gezielter Übung besprechen')
# Save independent student packets; the collection contains the same three packets.
for i,spec in enumerate(variants,1):
 d=new('Lernerfolgskontrolle Wunschbrief · Termin '+str(i));packet(d,spec,True);d.save(EX/f'Lernerfolgskontrolle_Wunschbrief_Termin_{i}.docx')
d=new('Lernerfolgskontrolle Wunschbrief · 3 Termine')
for i,spec in enumerate(variants):packet(d,spec,i==0)
d.save(EX/'Lernerfolgskontrolle_Wunschbrief_3_Termine.docx')
d=new('Probearbeit Wunschbrief · Komplettpaket');packet(d,mock,True);feedback(d);d.save(MOCK/'Probearbeit_Wunschbrief_Komplettpaket.docx')
print('Generated three individual packets, collection, and mock with feedback.')
