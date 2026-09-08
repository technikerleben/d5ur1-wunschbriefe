"""Optional paper projects and deeper reasoning challenges; classroom edition v3."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'materialien/projekte'
def new(title):
 d=Document();s=d.sections[0];s.page_width=Cm(14.8);s.page_height=Cm(21)
 s.top_margin=s.bottom_margin=Cm(.9);s.left_margin=s.right_margin=Cm(.9)
 for name,size in [('Normal',14),('Title',20),('Heading 1',16),('Heading 2',14)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('24313A')
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.space_before=Pt(0);st.paragraph_format.line_spacing=1.02;st.paragraph_format.keep_with_next=name!='Normal'
 d.core_properties.title=title+' · Fassung 3';d.core_properties.author='Deutschunterricht';d.core_properties.subject='Schritt 5 · Papiermaterial'
 return d
def p(d,t,style=None):return d.add_paragraph(t,style)
def band(d,label,t):
 x=p(d,'');x.add_run(label+'\n').bold=True;x.add_run(t);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF2F4');x._p.get_or_add_pPr().append(sh)
def page(d,label,title,first=False):
 x=p(d,label,'Title');x.paragraph_format.page_break_before=not first;p(d,title,'Heading 1')
def phase(d,label,t):p(d,label,'Heading 2');p(d,t)
d=new('Freiwillige Projekte')
page(d,'Freiwillige Projekte','Wähle, was dich interessiert',True)
p(d,'Nach der Probearbeit kannst du ein Projekt wählen. Es ist freiwillig. Es ersetzt weder die Probearbeit noch die Arbeit.')
band(d,'Dein Weg','Input → Übung → Vertiefung → Probearbeit → Projekte → Lernerfolgskontrolle (Arbeit)')
for t in ['1  Ein echter Wunschbrief','2  Eine Umfrage prüfen und nutzen','3  Eine Idee für einen Schulort','4  Ein verständliches Wunschposter','5  Zwei Briefe, zwei Empfänger','6  Auf eine Antwort reagieren']:p(d,t)
p(d,'Besprich deine Wahl mit der Lehrkraft. Du kannst allein oder nach Absprache mit einem Partnerkind arbeiten.')
band(d,'Du brauchst Papier','Arbeite im Heft oder auf Papier. Alle Projekte gehen ohne Laptop. Für Projekt 2 und 4 liegt eine Datenkarte bei.')
p(d,'Hilfen sind erlaubt. Die besonderen Projektaufträge gelten nur für dein gewähltes Projekt. Sie sind keine neuen Pflichtanforderungen für die Arbeit.')
projects=[
('Ein echter Wunschbrief','Ein Brief mit einem begründeten Vorschlag.','Heft, Briefpapier, Merkblätter 2 und 4.',
'Wähle einen Schulwunsch und eine passende Person. Sammle zwei Gründe. Entscheide, welcher für diese Person wichtiger ist.',
'Schreibe einen vollständigen Brief. Erkläre deinen wichtigsten Grund. Ergänze eine konkrete Bitte: Was soll die Person als Nächstes tun?',
'Zeige deinen stärksten Grund. Erkläre mündlich, warum du ihn gewählt hast. Prüfe den Brief mit Merkblatt 6.',
'Dein Wunsch ist klar. Dein Grund passt zur Person. Deine Bitte nennt einen machbaren nächsten Schritt.',
'Gib den Brief nur nach Absprache mit der Lehrkraft ab. Ein Übungsbrief ist auch möglich.'),
('Eine Umfrage prüfen und nutzen','Zwei Sätze zu Zahlen und ein kurzer Wunschbrief.','Heft, Datenkarte, Merkblatt 4.',
'Lies die Datenkarte. Wähle einen der vier Wünsche. Prüfe: Wer wurde befragt? Wie oft durfte jedes Kind wählen?',
'Schreibe zwei richtige Sätze zu den Zahlen. Schreibe dann einen vollständigen Brief zu deinem Wunsch. Nutze darin einen Satz zu den Zahlen und einen sachlichen Grund.',
'Erkläre: Was zeigen die Zahlen? Was sagen sie nicht über die ganze Schule?',
'Deine Zahlen stimmen. Dein Grund erklärt einen Nutzen. Du behauptest nicht, dass die ganze Schule befragt wurde.',
'Beginne: „In der Beispielklasse …“ Die Zahlen sind erfunden, keine echte Klassenumfrage.'),
('Eine Idee für einen Schulort','Eine Projektseite mit Skizze und kurzem Brief.','Papier, Stifte, Merkblätter 2 und 4.',
'Wähle einen Ort an der Schule. Beschreibe ein Problem. Sammle zwei mögliche Lösungen.',
'Vergleiche die Lösungen: Was hilft? Was könnte schwierig sein? Wähle eine Lösung. Zeichne sie und beschrifte sie. Schreibe dazu einen vollständigen Wunschbrief.',
'Nenne ein Bedenken gegen deine Wahl. Erkläre, wie du damit umgehen möchtest.',
'Ort und Problem sind klar. Du erklärst deine Wahl. Deine Skizze und dein Brief zeigen dieselbe Idee.',
'Du kannst eine kurze Probephase vorschlagen. Verlasse den Raum nur nach Absprache.'),
('Ein verständliches Wunschposter','Ein Poster aus Papier für eine kleine Ausstellung.','A3-Papier, Stifte, Datenkarte.',
'Wähle einen Wunsch aus der Datenkarte. Plane Platz für Überschrift, Zahl, Erklärung und eine freundliche Bitte.',
'Gestalte dein Poster. Nenne die richtige Zahl und die Beispielklasse. Ergänze einen Nutzen deines Wunsches. Die Bitte soll zeigen, was du erreichen möchtest.',
'Zeige das Poster einem Partnerkind. Frage: „Was wünsche ich mir? Was zeigt die Zahl?“ Verbessere nur unklare Stellen.',
'Dein Wunsch ist schnell erkennbar. Die Zahl passt dazu. Die Schrift ist lesbar. Die Daten sind als Beispiel markiert.',
'Ein Diagramm ist möglich, aber nicht nötig. Eine klar erklärte Zahl reicht.'),
('Zwei Briefe, zwei Empfänger','Zwei kurze Briefe mit demselben Wunsch.','Heft, Merkblätter 2 und 3.',
'Wähle einen Wunsch. Schreibe an ein Kind aus der Schülervertretung und an die Schulleitung. Überlege, wie beide helfen könnten.',
'Schreibe zwei vollständige Briefe. Passe Anrede, Erklärung und konkrete Bitte an die Person an. Tausche nicht nur die Namen aus.',
'Markiere zwei wichtige Unterschiede. Erkläre mündlich, warum du sie gewählt hast.',
'Der Wunsch bleibt gleich. Beide Personen werden passend angesprochen. Ihre unterschiedlichen Aufgaben werden deutlich.',
'Die Schülervertretung kann Wünsche sammeln. Die Schulleitung kann Möglichkeiten prüfen. Zusagen darfst du nicht erfinden.'),
('Auf eine Antwort reagieren','Ein Antwortbrief mit einem überarbeiteten Vorschlag.','Heft, Merkblätter 2 und 3.',
'Lies die Antwort: „Eine Spieleausleihe ist interessant. Uns fehlen aber Geld und Personen für die Ausgabe. Wie könnte ein kleiner Versuch aussehen?“',
'Plane einen Versuch mit vorhandenen Spielen. Wer könnte freiwillig helfen? Schreibe einen vollständigen Antwortbrief. Bedanke dich und gehe auf beide Schwierigkeiten ein.',
'Prüfe: Was müsste noch abgesprochen werden? Kennzeichne Ideen als Vorschlag, nicht als feste Zusage.',
'Du gehst auf Geld und Ausgabe ein. Dein Versuch ist konkret. Du versprichst nichts für andere Personen.',
'Beginne: „Vielen Dank für Ihre Antwort. Ich schlage vor, …“ Die Antwort oben ist ein erfundenes Beispiel.')]
for i,(title,product,material,plan,do,reflect,check,help_) in enumerate(projects,1):
 page(d,f'Projekt {i} · freiwillig',title)
 band(d,'Dein Ergebnis',product)
 p(d,'Material: '+material)
 for label,t in [('Planung',plan),('Durchführung',do),('Reflexion',reflect)]:phase(d,label,t)
 band(d,'Prüfe',check);p(d,'Hilfe: '+help_)
page(d,'Datenkarte','Für Projekt 2 und 4')
band(d,'Erfundene Beispielumfrage','26 Kinder einer Beispielklasse wurden gefragt: „Welchen Wunsch sollen wir zuerst besprechen?“ Jedes Kind durfte genau einen Wunsch wählen.')
for t in ['Spieleausleihe: 9 Stimmen','Leseecke: 7 Stimmen','Sitzbänke: 6 Stimmen','Neue Bälle: 4 Stimmen']:p(d,t)
p(d,'Zusammen: 26 Stimmen.','Heading 2')
band(d,'Das kannst du sagen','9 von 26 Kindern der Beispielklasse haben die Spieleausleihe gewählt. Sie hat unter diesen vier Wünschen die meisten Stimmen.')
band(d,'Das weißt du nicht','Die Umfrage zeigt nicht, was alle Kinder der Schule möchten. Sie zeigt auch nicht, welche anderen Wünsche ein Kind noch gut findet.')
p(d,'Du darfst auch einen Wunsch mit weniger Stimmen wählen. Begründe, welchen Nutzen er hat.')
p(d,'Echte Klassendaten kannst du nach Absprache stattdessen nutzen. Notiere dann die Frage, die Zahl der Befragten und ob mehrere Antworten erlaubt waren.')
d.save(OUT/'Freiwillige_Projekte_Wunschbriefe_DIN_A5.docx')

d=new('Freiwillige Vertiefungen')
vs=[('Gründe vergleichen','Du wählst einen starken Grund für eine Person.',
'Wunsch: eine Leseecke.\nA: „Ich finde Lesen toll.“\nB: „Dort können Kinder in der Pause ruhig lesen.“\nC: „Leseecken sind eben gut.“',
'Wähle eine Person. Ordne die Gründe von überzeugend bis wenig überzeugend. Erkläre deine stärkste Wahl. Verbessere einen schwachen Grund.',
'Würde ein Kind denselben Grund besonders wichtig finden wie die Schulleitung? Erkläre eine mögliche andere Sicht.',
'Dein Grund nennt einen Nutzen. Deine Erklärung passt zur Person. Eine andere Reihenfolge ist mit passender Erklärung möglich.'),
('Mit einem Bedenken umgehen','Du findest einen Vorschlag, der ein Problem ernst nimmt.',
'Wunsch: neue Sitzbänke auf dem Schulhof.\nBedenken: „Dort brauchen wir Platz zum Spielen.“',
'Sammle zwei Lösungen, die Sitzplätze und Spielfläche berücksichtigen. Vergleiche sie. Wähle eine Lösung und schreibe drei Sätze: Wunsch, Bedenken, Vorschlag.',
'Was müsste vor der Entscheidung noch geprüft werden? Formuliere eine konkrete Frage.',
'Du gehst auf den Platz zum Spielen ein. Dein Vorschlag hilft beim genannten Problem. Du behauptest nicht, dass schon alles geklärt ist.'),
('Eine Zahl kritisch lesen','Du erkennst, was eine Umfrage wirklich zeigt.',
'26 Kinder wählen jeweils genau einen Wunsch. 9 wählen die Spieleausleihe. Die anderen Wünsche erhalten 7, 6 und 4 Stimmen.\nBehauptung: „Fast alle Kinder unserer Schule wollen die Spieleausleihe.“',
'Erkläre zwei Fehler in der Behauptung. Schreibe einen richtigen Satz zu den Zahlen. Formuliere eine Frage, die diese Umfrage noch nicht beantwortet.',
'Kann der Wunsch trotzdem sinnvoll sein? Ergänze einen sachlichen Grund, der ohne die Zahl verständlich ist.',
'Du beachtest 9 von 26 und die befragte Gruppe. Du unterscheidest Zahlen und sachlichen Nutzen.'),
('Eine faire Entscheidung finden','Du berücksichtigst unterschiedliche Bedürfnisse.',
'Zwei Ideen für die Pause: eine ruhige Leseecke und mehr Platz für Ballspiele. Manche Kinder möchten Ruhe. Andere möchten sich bewegen.',
'Beschreibe für jede Idee einen Nutzen. Entwickle einen gemeinsamen Vorschlag. Schreibe eine freundliche Bitte an die Schulleitung mit deiner Lösung.',
'Nenne eine Grenze deiner Lösung. Schlage vor, wie ihr herausfinden könnt, ob sie im Alltag hilft.',
'Du nimmst beide Bedürfnisse ernst. Dein Vorschlag ist konkret. Ein kleiner Versuch mit anschließender Rückmeldung ist möglich.')]
for i,(title,goal,example,task,extra,check) in enumerate(vs,1):
 page(d,f'Vertiefung {i} · freiwillig',title,i==1)
 band(d,'Dein Ziel',goal);band(d,'Dein Ausgangspunkt',example)
 phase(d,'Planung', 'Lies das Beispiel. Sage zuerst mündlich, welche Ideen du hast.')
 phase(d,'Durchführung · Im Heft',task)
 phase(d,'Noch weiter · freiwillig',extra)
 phase(d,'Reflexion',check)
 p(d,'Hilfe: Nutze Merkblatt 4. Besprich deine Ideen zuerst mündlich.')
d.save(OUT/'Freiwillige_Vertiefungen_Wunschbriefe_DIN_A5.docx')
print('Generated project package (8 A5 pages) and four optional challenges.')
